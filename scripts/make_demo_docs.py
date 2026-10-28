"""Regenerate the PDF and DOCX files of the synthetic demo corpus.

Usage: .venv/bin/python scripts/make_demo_docs.py
"""
from pathlib import Path

import docx
import fitz

OUT = Path(__file__).resolve().parent.parent / "demo_corpus"
BANNER = "SYNTHETIC DOCUMENT. Northbridge Institute of Technology is fictional; this file exists only to test Sentinel-RAG."


def make_pdf():
    pages = [
        ("Laboratory Safety Policy", [
            BANNER,
            "Personal Protective Equipment",
            "Safety goggles and a lab coat must be worn at all times in chemistry laboratories. "
            "Closed-toe shoes are mandatory in every laboratory.",
        ]),
        ("Incident Reporting", [
            "Every laboratory incident, including minor spills, must be reported to the lab supervisor "
            "within 24 hours using the incident form.",
            "Fume Hoods",
            "Work involving volatile solvents must be carried out inside a fume hood.",
        ]),
    ]
    doc = fitz.open()
    for title, paras in pages:
        page = doc.new_page()
        y = 72
        page.insert_text((72, y), title, fontsize=16)
        y += 36
        for p in paras:
            rect = fitz.Rect(72, y, 523, y + 120)
            page.insert_textbox(rect, p, fontsize=11)
            y += 30 + 14 * (len(p) // 80 + 1)
    doc.save(OUT / "lab_safety_policy.pdf")


def make_docx():
    d = docx.Document()
    d.add_heading("Merit Scholarship Policy", 0)
    d.add_paragraph(BANNER)
    d.add_heading("Merit Scholarship", 1)
    d.add_paragraph("The merit scholarship covers 50 percent of the tuition fee.")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Criterion", "Requirement"),
                                ("Initial award", "CGPA of 8.5 or above in the qualifying year"),
                                ("Renewal", "CGPA of 8.0 or above in each academic year")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_heading("Medical Break", 1)
    d.add_paragraph("A student on an approved medical break retains the scholarship for one semester "
                    "without meeting the renewal CGPA.")
    d.save(OUT / "scholarship_policy.docx")


if __name__ == "__main__":
    make_pdf()
    make_docx()
    print("wrote", OUT / "lab_safety_policy.pdf", OUT / "scholarship_policy.docx")
