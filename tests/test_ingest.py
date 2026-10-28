import io

import docx
import fitz
import pytest

from sentinel_rag import models
from sentinel_rag.config import Settings
from sentinel_rag.ingest import IngestError, chunk_blocks, parse


def make_pdf(pages, **save):
    doc = fitz.open()
    for text in pages:
        p = doc.new_page()
        if text:
            p.insert_textbox(fitz.Rect(72, 72, 520, 700), text, fontsize=11)
    buf = io.BytesIO()
    doc.save(buf, **save)
    return buf.getvalue()


def test_pdf_keeps_page_numbers():
    data = make_pdf(["Alpha rule applies to every student in the first year.",
                     "Beta rule applies only to final year students."])
    parsed = parse("x.pdf", data)
    assert [b.page for b in parsed.blocks] == [1, 2]
    chunks = chunk_blocks(parsed.blocks, "x", 1, Settings())
    assert chunks[0].page_start == 1 and chunks[-1].page_end == 2
    assert "p. 1" in chunks[0].locator()


def test_pdf_empty_page_is_reported_not_silently_dropped():
    parsed = parse("x.pdf", make_pdf(["Some real text on the first page here.", ""]))
    assert any("page 2" in w for w in parsed.warnings)


def test_encrypted_pdf_is_a_parser_error():
    data = make_pdf(["secret text"], encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(IngestError, match="encrypted"):
        parse("x.pdf", data)


def test_docx_tables_keep_column_associations_and_sections():
    d = docx.Document()
    d.add_heading("Fees", 1)
    d.add_paragraph("Fees are payable each semester.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Amount"
    t.cell(1, 0).text, t.cell(1, 1).text = "Library", "INR 200"
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse("f.docx", buf.getvalue())
    texts = [b.text for b in parsed.blocks]
    assert "Item: Library; Amount: INR 200." in texts
    assert all(b.section == "Fees" for b in parsed.blocks)


def test_text_headings_become_sections_and_numbers_survive_normalisation():
    raw = b"# Rules\n\nThe fee is INR 1,500 and is not refundable.\nIt is due   within 14 days.\n"
    parsed = parse("r.md", raw)
    assert parsed.blocks[0].section == "Rules"
    assert "INR 1,500" in parsed.blocks[0].text and "not refundable" in parsed.blocks[0].text
    assert "  " not in parsed.blocks[0].text


def test_chunks_never_exceed_embedder_limit_and_overlap_is_bounded():
    s = Settings()
    sentence = "Students must submit the signed form to the department office before the deadline. "
    parsed = parse("long.txt", (sentence * 80).encode())
    chunks = chunk_blocks(parsed.blocks, "long", 1, s)
    assert len(chunks) > 3
    for c in chunks:
        assert c.n_tokens <= s.chunk_tokens
        assert len(models.tokenizer().tokenize(c.embed_text())) + 2 <= models.max_seq_length()
        assert c.overlap_tokens <= s.overlap_tokens
    assert any(c.overlap_tokens > 0 for c in chunks[1:])


def test_unsupported_and_empty_files_fail_clearly():
    with pytest.raises(IngestError, match="unsupported"):
        parse("x.png", b"123")
    with pytest.raises(IngestError, match="empty"):
        parse("x.txt", b"")
